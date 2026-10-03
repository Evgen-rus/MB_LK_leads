from datetime import datetime
import json
from pathlib import Path

import pandas as pd
import pytest
import uuid
from io import BytesIO
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import models
from backend.app.lead_analytics import db
from backend.app.lead_analytics import export_history, router as analytics_router
from backend.app.lead_analytics.export_history import ExportMetadata, ExportPeriod, save_analysis_export
from backend.app.lead_analytics.models import ColumnMapping, StatusRule
from backend.app.lead_analytics.router import _export_lk_snapshot, build_router


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "analytics.db")
    db.init_db()
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def get_db():
        session = Session()
        try:
            yield session
        finally:
            session.close()

    app_state_session = Session

    def admin_dependency():
        return {"id": 1}

    app = FastAPI()
    app.state.test_session_factory = app_state_session
    app.include_router(build_router(admin_dependency, get_db, lambda v: v or "", lambda v: v or "", lambda v: v or ""))
    with Session() as session:
        session.add_all([
            models.User(id=1, login="root", password_hash="x", role="admin"),
            models.User(id=2, login="client", password_hash="x", role="client"),
            models.User(id=3, login="another", password_hash="x", role="client"),
            models.User(id=4, login="agent", password_hash="x", role="agent"),
        ])
        session.add(models.ClientProfile(user_id=2, name="Client", work_status="В работе"))
        session.add_all([
            models.Project(id=20, user_id=2, name="Active Pixel", tag="x", collection_source="Звонки",
                           data_source_code="UNMAPPED", status="Активен", delivery_status="Активна", data_limit=0),
            models.Project(id=21, user_id=2, name="Deleted Provider", tag="x", collection_source="Сайты",
                           data_source_code="B1", status="Удалён", delivery_status="Отключена", data_limit=0,
                           deleted_at=datetime(2026, 1, 1)),
            models.Project(id=30, user_id=3, name="Other", tag="x", collection_source="Сайты",
                           data_source_code="B1", status="Активен", delivery_status="Активна", data_limit=0),
        ])
        session.commit()
    return TestClient(app)


def test_clients_projects_and_independent_group_membership(api):
    response = api.get("/admin/analytics/clients")
    assert response.status_code == 200, response.text
    assert {item["id"] for item in response.json()} == {2, 3}
    projects = api.get("/admin/analytics/clients/2/projects?q=Pixel").json()
    assert projects == [{"id": 20, "name": "Active Pixel", "status": "Активен", "deleted_at": None,
                         "collection_source": "Звонки"}]
    all_projects = api.get("/admin/analytics/clients/2/projects").json()
    assert {item["id"] for item in all_projects} == {20, 21}
    assert next(item for item in all_projects if item["id"] == 21)["deleted_at"]

    created = api.post("/admin/analytics/clients/2/groups", json={
        "name": "Combined", "project_ids": [20, 21], "spreadsheet_url": None,
    })
    assert created.status_code == 200
    group = created.json()
    assert group["project_ids"] == [20, 21]
    assert api.get(f"/admin/analytics/groups/{group['id']}/status-rules").json()["system_rules"] == []
    db.ensure_project(db.group_key(group["id"]))
    db.add_status_rule(StatusRule(pattern="New", match_type="exact", group_name="Качественные",
                                  project_code=db.group_key(group["id"])))
    same_name = api.post("/admin/analytics/clients/2/groups", json={
        "name": "Combined", "project_ids": [20], "spreadsheet_url": None,
    }).json()
    assert api.get(f"/admin/analytics/groups/{same_name['id']}/status-rules").json()["project_rules"] == []
    assert api.put(f"/admin/analytics/groups/{group['id']}", json={
        "name": "Renamed", "project_ids": [21], "spreadsheet_url": None,
    }).json()["id"] == group["id"]
    assert api.get(f"/admin/analytics/groups/{group['id']}/status-rules").json()["project_rules"][0]["pattern"] == "New"
    assert api.post("/admin/analytics/clients/2/groups", json={
        "name": "Bad", "project_ids": [30], "spreadsheet_url": None,
    }).status_code == 400


def test_router_requires_admin_dependency():
    app = FastAPI()

    def denied():
        raise HTTPException(status_code=403)

    app.include_router(build_router(denied, lambda: None, lambda v: v or "", lambda v: v or "", lambda v: v or ""))
    assert TestClient(app).get("/admin/analytics/clients").status_code == 403


def test_internal_lk_snapshot_reads_all_rows_over_old_export_limit(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    with Session() as session:
        session.add_all(models.ProviderLead(
            vid=str(index), project_id=70, imported_at=datetime(2026, 1, 2), phone="+70000000000"
        ) for index in range(15_001))
        session.commit()
        target = tmp_path / "lk.xlsx"
        count = _export_lk_snapshot(session, 2, [70], datetime(2026, 1, 1), datetime(2026, 1, 3),
                                    target, lambda v: v or "", lambda v: v or "", lambda v: v or "")
    from openpyxl import load_workbook

    assert count == 15_001
    workbook = load_workbook(target, read_only=True)
    assert sum(1 for _ in workbook.active.iter_rows()) == 15_002
    workbook.close()


def test_internal_lk_snapshot_cleans_openpyxl_temp_file_on_iterator_error(tmp_path, monkeypatch):
    from openpyxl.worksheet import _writer as openpyxl_writer

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    models.Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    created_temp_files = []
    original_create_temp_file = openpyxl_writer.create_temporary_file

    def track_temp_file(*args, **kwargs):
        path = original_create_temp_file(*args, **kwargs)
        created_temp_files.append(path)
        return path

    monkeypatch.setattr(openpyxl_writer, "create_temporary_file", track_temp_file)

    def broken_export(*_args, **_kwargs):
        yield {
            "imported_at": datetime(2026, 1, 2), "phone": "+70000000000", "source": "source",
            "utm_campaign": "campaign", "project_name": "project", "lk_id": "lkid",
            "ext_id": "external", "user_name": "client",
        }
        raise RuntimeError("synthetic iterator failure")

    monkeypatch.setattr(analytics_router.crud, "iter_provider_leads_for_export", broken_export)
    target = tmp_path / "failed.xlsx"
    with Session() as session, pytest.raises(RuntimeError, match="synthetic iterator failure"):
        _export_lk_snapshot(session, 2, [70], datetime(2026, 1, 1), datetime(2026, 1, 3),
                            target, lambda value: value or "", lambda value: value or "",
                            lambda value: value or "")

    assert created_temp_files
    assert not any(Path(name).exists() for name in created_temp_files)
    assert not (set(created_temp_files) & set(openpyxl_writer.ALL_TEMP_FILES))
    assert not target.exists()


def _add_run(group_id: int, run_id: str, group_name: str = "Original group"):
    group = db.get_group(group_id)
    db.create_run({
        "id": run_id, "group_id": group_id, "client_id": group["client_id"], "group_name": group_name,
        "project_ids": [20], "project_names": {"20": "Active Pixel"},
        "periods": [{"period_start": "2026-01-01", "period_end": "2026-01-31"}],
        "settings": {}, "lk_snapshot": "input/lk.xlsx", "client_snapshot": "input/client.xlsx",
        "source_file_name": "client.xlsx",
    })


def test_analyze_setup_restores_saved_mapping_and_sheet(api, tmp_path, monkeypatch):
    runs_dir = tmp_path / "runs"
    monkeypatch.setattr(analytics_router, "RUNS_DIR", runs_dir)
    group_id = db.create_group(2, "Saved setup", [20], None)
    run_id = "saved-column-mapping"
    _add_run(group_id, run_id)
    output_dir = runs_dir / run_id / "output"
    output_dir.mkdir(parents=True)
    workbook = Workbook()
    first = workbook.active
    first.title = "Сопоставленные"
    first.append(["Created", "Client status", "Phone"])
    first.append(["2026-01-03", "New lead", "+70000000000"])
    selected = workbook.create_sheet("Other")
    selected.append(["Дата", "Статус"])
    selected.append(["2026-01-03", "irrelevant"])
    workbook.save(output_dir / "joined_сопоставление.xlsx")

    key = db.group_key(group_id)
    db.ensure_project(key)
    db.save_column_mapping(key, ColumnMapping(
        sheet_name="Сопоставленные", date_column="Created", phone_column="Phone", status_column="Client status",
    ))
    response = api.post(f"/admin/analytics/groups/{group_id}/runs/{run_id}/analyze/setup", json={})

    assert response.status_code == 200, response.text
    assert response.json()["mapping"]["sheet_name"] == "Сопоставленные"
    assert response.json()["mapping"]["date_column"] == "Created"
    assert response.json()["mapping"]["status_column"] == "Client status"
    assert response.json()["unknown_statuses"] == ["New lead"]


def test_completed_analyze_job_exposes_its_exact_export_id(api, tmp_path, monkeypatch):
    runs_dir = tmp_path / "runs"
    monkeypatch.setattr(analytics_router, "RUNS_DIR", runs_dir)
    group_id = db.create_group(2, "Snapshot group", [20], None)
    run_id = "analysis-export-link"
    _add_run(group_id, run_id)
    run_dir = runs_dir / run_id
    output_dir = run_dir / "output"
    output_dir.mkdir(parents=True)
    (output_dir / "joined_сопоставление.xlsx").write_bytes(b"prepared matched workbook")
    key = db.group_key(group_id)
    db.ensure_project(key)

    def insert_export(export_number: int, linked_run: str) -> int:
        with db.connect() as conn:
            cursor = conn.execute(
                """INSERT INTO analysis_exports(project_code,export_number,period_start,period_end,analysis_date,
                source_file_name,run_id,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?)""",
                (key, export_number, "2026-01-01", "2026-01-31", "2026-02-01", "client.xlsx", linked_run,
                 db.now_text(), db.now_text()),
            )
            return int(cursor.lastrowid)

    created: list[int] = []

    def fake_analyze(_project, _match_path, _mapping, output_path, _metadata, **kwargs):
        output = Path(output_path) / "analysis.xlsx"
        output.write_bytes(b"generated workbook")
        created.append(insert_export(1, kwargs["run_id"]))
        newer = insert_export(2, "another-concurrent-run")
        assert newer > created[0]
        return output

    monkeypatch.setattr(analytics_router, "analyze_file", fake_analyze)
    queued = db.create_processing_job(run_id, "analyze", {
        "project": "Original group", "storage_key": key,
        "mapping": {"sheet_name": "joined", "status_column": "Status", "date_column": "Date"},
        "rule_snapshot": [], "periods": [{"period_start": "2026-01-01", "period_end": "2026-01-31"}],
        "analysis_date": None, "source_file_name": "client.xlsx", "match_output": "joined_сопоставление.xlsx",
        "settings_snapshot": {},
    }, group_id=group_id)

    analytics_router._run_job(queued)
    result = api.get(f"/admin/analytics/groups/{group_id}/jobs/{queued['id']}")
    assert result.status_code == 200
    assert result.json()["status"] == "completed", result.json().get("error_text")
    assert result.json()["export_id"] == created[0]
    assert result.json()["export_id"] != db.get_analysis_export_id(key, "another-concurrent-run")
    assert result.json()["output_file_name"].startswith(f"job-{queued['id']}_")


def test_archived_group_keeps_history_and_deletes_only_selected_report(api, tmp_path, monkeypatch):
    report_dir = tmp_path / "reports"
    report_dir.mkdir()
    monkeypatch.setattr(export_history, "ANALYSIS_REPORTS_DIR", report_dir)
    group_id = db.create_group(2, "Reports", [20], None)
    key = db.group_key(group_id)
    period = ExportPeriod("2026-01-01", "2026-01-31")
    totals = pd.DataFrame({
        "Всего идентификаций": [1], "Недозвон": [0], "Недозвон %": [0.0],
        "Качественные": [1], "Кач. %": [1.0], "Сигнал спроса": [1], "Сигнал спроса %": [1.0],
    })
    saved_ids = []
    report_names = []
    for number in (1, 2):
        filename = f"{uuid.uuid4().hex}.xlsx"
        report_names.append(filename)
        (report_dir / filename).write_bytes(f"workbook-{number}".encode())
        saved_ids.append(save_analysis_export(
            key, ExportMetadata(None, [period]),
            [(period, totals, {"domain_channel": pd.DataFrame(), "source_channel": pd.DataFrame(), "channel": pd.DataFrame()})],
            report_file_name=filename, run_id=f"run-{number}",
        ))

    assert api.delete(f"/admin/analytics/groups/{group_id}").status_code == 200
    history = api.get(f"/admin/analytics/groups/{group_id}/exports")
    assert history.status_code == 200
    assert [item["id"] for item in history.json()] == saved_ids
    downloaded = api.get(f"/admin/analytics/groups/{group_id}/exports/{saved_ids[0]}/download")
    assert downloaded.status_code == 200
    assert downloaded.content == b"workbook-1"
    assert api.delete(f"/admin/analytics/groups/{group_id}/exports/{saved_ids[0]}").status_code == 200
    assert [item["id"] for item in api.get(f"/admin/analytics/groups/{group_id}/exports").json()] == [saved_ids[1]]
    assert not (report_dir / report_names[0]).exists()
    assert (report_dir / report_names[1]).exists()


def test_invalid_client_xlsx_does_not_create_run(api, tmp_path, monkeypatch):
    runs_dir = tmp_path / "runs"
    monkeypatch.setattr(analytics_router, "RUNS_DIR", runs_dir)
    group_id = api.post("/admin/analytics/clients/2/groups", json={
        "name": "Upload test", "project_ids": [20], "spreadsheet_url": None,
    }).json()["id"]
    response = api.post(
        f"/admin/analytics/groups/{group_id}/runs",
        data={"periods": '[{"period_start":"2026-01-01","period_end":"2026-01-31"}]'},
        files={"client_file": ("client.xlsx", b"not an xlsx workbook", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 400
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM analytics_runs").fetchone()[0] == 0
    assert not list(runs_dir.glob("*/"))


def test_empty_group_never_exports_every_client_project(api, monkeypatch):
    group_id = db.create_group(2, "Empty group", [], None)
    monkeypatch.setattr(analytics_router, "_export_lk_snapshot", lambda *args, **kwargs: pytest.fail("must reject empty group"))
    response = api.post(
        f"/admin/analytics/groups/{group_id}/runs",
        data={"periods": '[{"period_start":"2026-01-01","period_end":"2026-01-31"}]'},
        files={"client_file": ("client.xlsx", b"unused", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 400
    assert "проекты" in response.json()["detail"].lower()


def test_google_access_error_is_sanitized_and_creates_no_run(api, tmp_path, monkeypatch):
    runs_dir = tmp_path / "runs"
    monkeypatch.setattr(analytics_router, "RUNS_DIR", runs_dir)
    group_id = db.create_group(2, "Sheets", [20], "https://docs.google.com/spreadsheets/d/synthetic-id/edit")
    monkeypatch.setattr(analytics_router, "read_spreadsheet_url",
                        lambda _url: (_ for _ in ()).throw(RuntimeError("private credential detail")))
    response = api.post(f"/admin/analytics/groups/{group_id}/runs", data={
        "periods": '[{"period_start":"2026-01-01","period_end":"2026-01-31"}]',
    })
    assert response.status_code == 502
    assert "private credential detail" not in response.text
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM analytics_runs").fetchone()[0] == 0
    assert not runs_dir.exists()


def test_prepare_freezes_multiple_periods_membership_and_includes_last_day(api, tmp_path, monkeypatch):
    runs_dir = tmp_path / "runs"
    monkeypatch.setattr(analytics_router, "RUNS_DIR", runs_dir)
    session_factory = api.app.state.test_session_factory
    with session_factory() as session:
        session.add_all([
            models.ProviderLead(vid="end-day", project_id=20, imported_at=datetime(2026, 1, 31, 23, 59, 59), phone="+70000000001"),
            models.ProviderLead(vid="next-day", project_id=20, imported_at=datetime(2026, 2, 1, 0, 0, 0), phone="+70000000002"),
        ])
        session.commit()
    group_id = db.create_group(2, "Before rename", [20], None)
    client_bytes = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Дата", "Телефон", "Статус"])
    sheet.append(["2026-01-31", "+70000000001", "New"])
    workbook.save(client_bytes)
    periods = [
        {"period_start": "2026-01-01", "period_end": "2026-01-15"},
        {"period_start": "2026-01-16", "period_end": "2026-01-31"},
    ]
    response = api.post(
        f"/admin/analytics/groups/{group_id}/runs",
        data={"periods": json.dumps(periods)},
        files={"client_file": ("client.xlsx", client_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    prepared = response.json()
    assert prepared["lk_row_count"] == 1
    assert [item["period_start"] for item in db.get_run(prepared["run_id"])["periods"]] == ["2026-01-01", "2026-01-16"]
    assert db.get_run(prepared["run_id"])["group_name"] == "Before rename"
    assert db.get_run(prepared["run_id"])["project_ids"] == [20]
    assert db.update_group(group_id, "After rename", [21], None)
    frozen = db.get_run(prepared["run_id"])
    assert frozen["group_name"] == "Before rename"
    assert frozen["project_ids"] == [20]
    assert frozen["project_names"] == {"20": "Active Pixel"}
