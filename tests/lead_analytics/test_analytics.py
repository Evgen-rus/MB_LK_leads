import pandas as pd
from openpyxl import load_workbook

from backend.app.lead_analytics import pipeline
from backend.app.lead_analytics.analytics import summarize
from backend.app.lead_analytics.export_history import ExportMetadata, ExportPeriod
from backend.app.lead_analytics.models import ColumnMapping
from backend.app.lead_analytics.report_writer import write_excel


def export_metadata(*periods: tuple[str, str]) -> ExportMetadata:
    return ExportMetadata(
        export_number=None,
        periods=[ExportPeriod(start, end) for start, end in periods],
        source_file_name="input.xlsx",
    )


def test_summarize_excludes_do_not_count_group():
    df = pd.DataFrame(
        {
            "Канал": ["A", "A", "A", "A"],
            "Группа статуса": ["Качественные", "Недозвон", "Не учитывать", "Еще не звонили"],
            "Исходный статус": ["В работе", "Недозвон", "", "Новый"],
        }
    )
    result = summarize(df, ["Канал"])
    assert result.loc[0, "Всего идентификаций"] == 2
    assert result.loc[0, "Сигнал спроса"] == 1
    assert result.loc[0, "Сигнал спроса %"] == 0.5
    assert result.loc[0, "Не обработано"] == 2
    assert result.loc[0, "Не обработано %"] == 0.5
    missed_index = list(result.columns).index("Недозвон %")
    assert list(result.columns)[missed_index + 1 : missed_index + 3] == ["Не обработано", "Не обработано %"]
    assert "Еще не звонили" not in result.columns


def test_summarize_places_competitor_after_demand_signal_and_excludes_it_from_demand():
    df = pd.DataFrame(
        {
            "Группа статуса": [
                "Качественные",
                "Уже наши / уже купил",
                "Конкурент",
                "Недозвон",
            ],
        }
    )
    result = summarize(df, [])
    columns = list(result.columns)
    already_percent_index = columns.index("Уже наши / купил %")
    assert columns[already_percent_index + 1] == "Сигнал спроса"
    assert columns[already_percent_index + 2] == "Сигнал спроса %"
    assert columns[already_percent_index + 3] == "Конкурент"
    assert columns[already_percent_index + 4] == "Конкурент %"
    assert result.loc[0, "Сигнал спроса"] == 2
    assert result.loc[0, "Сигнал спроса %"] == 0.5
    empty = pd.DataFrame(columns=["Канал", "Группа статуса"])
    assert list(summarize(empty, ["Канал"]).columns)[1:] == list(result.columns)


def test_analyze_file_loads_status_rules_once(monkeypatch, tmp_path):
    source = pd.DataFrame(
        {
            "Дата": ["2026-01-01", "2026-01-02"],
            "Телефон": ["1", "2"],
            "Источник": ["example.com", "example.com"],
            "Статус": ["Новый", "Новый"],
        }
    )
    mapping = ColumnMapping(
        sheet_name="Сопоставленные",
        date_column="Дата",
        phone_column="Телефон",
        source_column="Источник",
        status_column="Статус",
    )
    rules = [object()]
    calls = 0
    saved_export = {}

    def write_report(path, _):
        path.write_bytes(b"complete workbook")
        return path

    def save_export(*_args, **kwargs):
        saved_export.update(kwargs)
        return 1

    monkeypatch.setattr(pipeline, "read_excel_sheet", lambda *_: source)

    def load_rules(project):
        nonlocal calls
        calls += 1
        return rules

    def classify_with_rules(status, comment, project, loaded_rules):
        assert loaded_rules is rules
        return "Качественные", "тест"

    monkeypatch.setattr(pipeline, "sorted_rules", load_rules)
    monkeypatch.setattr(pipeline, "classify", classify_with_rules)
    monkeypatch.setattr(pipeline, "write_excel", write_report)
    monkeypatch.setattr(pipeline, "save_analysis_export", save_export)
    monkeypatch.setattr(pipeline, "ANALYSIS_REPORTS_DIR", tmp_path / "reports")

    pipeline.analyze_file(
        "p",
        "input.xlsx",
        mapping,
        tmp_path,
        export_metadata(("2026-01-01", "2026-01-02")),
    )

    assert calls == 1
    archived_report = tmp_path / "reports" / saved_export["report_file_name"]
    assert archived_report.read_bytes() == b"complete workbook"


def test_analyze_file_uses_matched_sheet_not_duplicates(tmp_path, monkeypatch):
    path = tmp_path / "match.xlsx"
    matched = pd.DataFrame(
        {
            "Дата": ["2026-01-01", "2026-01-02"],
            "Телефон": ["1", "2"],
            "Канал": ["A", "A"],
            "Источники": ["example.com", "example.com"],
            "Статус клиента": ["Новый", "Новый"],
            "Комментарий клиента": ["", ""],
        }
    )
    duplicates = pd.DataFrame(
        {
            "Дата создания": ["2026-01-01"] * 5,
            "Рабочий телефон": ["1"] * 5,
            "_status": ["СПАМ"] * 5,
            "_comment": [""] * 5,
            "Источник": ["x"] * 5,
        }
    )
    with pd.ExcelWriter(path) as writer:
        matched.to_excel(writer, sheet_name="Сопоставленные", index=False)
        duplicates.to_excel(writer, sheet_name="Дубли клиента", index=False)

    mapping = ColumnMapping(
        sheet_name="Дубли клиента",
        date_column="Дата",
        phone_column="Телефон",
        channel_column="Канал",
        source_column="Источники",
        status_column="Статус клиента",
        comment_column="Комментарий клиента",
    )
    captured = {}

    def fake_write(output, sheets):
        captured["data"] = next(frame for name, frame in sheets.items() if name.startswith("Данные "))
        output.write_bytes(b"workbook")
        return output

    monkeypatch.setattr(pipeline, "sorted_rules", lambda _project: [])
    monkeypatch.setattr(pipeline, "classify", lambda *_args, **_kwargs: ("Качественные", "тест"))
    monkeypatch.setattr(pipeline, "write_excel", fake_write)
    monkeypatch.setattr(pipeline, "save_analysis_export", lambda *_args, **_kwargs: 1)
    monkeypatch.setattr(pipeline, "ANALYSIS_REPORTS_DIR", tmp_path / "reports")

    pipeline.analyze_file(
        "p",
        path,
        mapping,
        tmp_path,
        export_metadata(("2026-01-01", "2026-01-02")),
    )

    assert len(captured["data"]) == 2
    assert list(captured["data"]["Телефон"]) == ["1", "2"]


def test_analyze_file_builds_five_sheets_for_each_period(tmp_path, monkeypatch):
    source = pd.DataFrame(
        {
            "Дата": ["2026-01-01", "2026-01-31", "2026-02-01", "нет даты"],
            "Канал": ["A", "A", "B", "B"],
            "Источник": ["one.ru", "one.ru", "two.ru", "two.ru"],
            "Статус": ["Новый"] * 4,
        }
    )
    mapping = ColumnMapping(
        sheet_name="Сопоставленные",
        date_column="Дата",
        channel_column="Канал",
        source_column="Источник",
        status_column="Статус",
    )
    captured = {}
    monkeypatch.setattr(pipeline, "read_excel_sheet", lambda *_: source)
    monkeypatch.setattr(pipeline, "sorted_rules", lambda _project: [])
    monkeypatch.setattr(pipeline, "classify", lambda *_args: ("Качественные", "тест"))

    def fake_write(path, sheets):
        captured.update(sheets)
        path.write_bytes(b"workbook")
        return path

    monkeypatch.setattr(pipeline, "write_excel", fake_write)
    monkeypatch.setattr(pipeline, "save_analysis_export", lambda *_args, **_kwargs: 1)
    monkeypatch.setattr(pipeline, "ANALYSIS_REPORTS_DIR", tmp_path / "reports")

    pipeline.analyze_file(
        "p",
        "input.xlsx",
        mapping,
        tmp_path,
        export_metadata(("2026-01-01", "2026-01-31"), ("2026-02-01", "2026-02-28")),
    )

    assert list(captured) == [
        "Итог",
        "По доменам 01.01-31.01",
        "По источникам 01.01-31.01",
        "По каналам 01.01-31.01",
        "Статусы 01.01-31.01",
        "Данные 01.01-31.01",
        "По доменам 01.02-28.02",
        "По источникам 01.02-28.02",
        "По каналам 01.02-28.02",
        "Статусы 01.02-28.02",
        "Данные 01.02-28.02",
    ]
    assert list(captured["Итог"]["Всего идентификаций"]) == [2, 1]
    assert len(captured["Данные 01.01-31.01"]) == 2
    assert "Выводы" not in captured


def test_write_excel_preserves_report_formatting(tmp_path):
    path = write_excel(
        tmp_path / "report.xlsx",
        {
            "Итог": pd.DataFrame(
                {"Кач. %": [0.03, 0.0910891089108911], "Канал": ["Поиск", "Сети"]}
            )
        },
    )

    ws = load_workbook(path).active

    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    assert ws["A2"].number_format == "0.00%"
    assert ws["A3"].number_format == "0.00%"
    assert ws["A2"].fill.fill_type == "solid"
    assert ws["A3"].fill.fill_type == "solid"
    assert ws.column_dimensions["B"].width >= len("Канал") + 2
