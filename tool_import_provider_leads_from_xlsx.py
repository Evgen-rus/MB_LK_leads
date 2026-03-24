"""
Импорт лидов провайдера из ручной XLSX-выгрузки.

Поддерживает два режима:
- preview (по умолчанию): только анализ файла и будущего импорта;
- import: реальная запись новых строк в provider_leads.

Запуск:
    python tool_import_provider_leads_from_xlsx.py --file "1774340403_report_434771_440089.xlsx" --preview
    python tool_import_provider_leads_from_xlsx.py --file "1774340403_report_434771_440089.xlsx" --preview --verbose
    python tool_import_provider_leads_from_xlsx.py --file "1774340403_report_434771_440089.xlsx" --import

Маппинг полей согласован с текущим webhook-потоком:
- id -> vid
- Проект -> project_name
- Телефон -> phone / phones_raw
- Создано -> prov_created_at
- Комментарий -> subdomain
- prov_chanel / prov_source вычисляются из project_name
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError

from backend.app import crud
from backend.app import db as db_mod
from backend.app import models
from backend.app.time_utils import now_msk


DEFAULT_DB_URL = "sqlite:///./app.db"
REQUIRED_HEADERS = ["id", "Проект", "Телефон", "Создано", "Комментарий"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Импорт лидов из XLSX-выгрузки провайдера в provider_leads"
    )
    parser.add_argument(
        "--file",
        required=True,
        help="Путь к XLSX-файлу выгрузки",
    )
    parser.add_argument(
        "--db-url",
        default=os.getenv("DATABASE_URL", DEFAULT_DB_URL),
        help="DATABASE_URL для подключения (по умолчанию из env или sqlite:///./app.db)",
    )
    parser.add_argument(
        "--preview",
        action="store_true",
        help="Только анализ файла без записи в БД (режим по умолчанию, если --import не указан)",
    )
    parser.add_argument(
        "--import",
        dest="do_import",
        action="store_true",
        help="Реально записать новые строки в provider_leads",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Обработать только первые N строк файла (0 = все строки)",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Печатать примеры проблемных строк",
    )
    return parser.parse_args()


def clean_scalar(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.lower() == "none":
        return None
    return text


def normalize_phone(value: Any) -> Optional[str]:
    text = clean_scalar(value)
    if not text:
        return None
    return "".join(text.split())


def parse_created_at(value: Any) -> Tuple[Optional[datetime], Optional[str]]:
    text = clean_scalar(value)
    if not text:
        return None, "empty_created_at"

    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt), None
        except ValueError:
            continue
    return None, f"invalid_created_at:{text}"


def parse_page_parts(page: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    """
    Держим логику синхронной с backend/app/main.py::_parse_page_parts.
    """
    if not page:
        return None, None
    parts = page.split("_")
    prov_chanel = parts[0].strip() if parts else None
    prov_source = None
    if len(parts) >= 3:
        prov_source = "_".join(parts[2:]).strip() or None
    return prov_chanel, prov_source


def validate_required_headers(headers: List[str]) -> None:
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        missing_str = ", ".join(missing)
        raise SystemExit(f"XLSX is missing required columns: {missing_str}")


def load_xlsx_rows(file_path: str, limit: int = 0) -> Tuple[List[str], List[Dict[str, Any]]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows_iter = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise SystemExit("XLSX file is empty.") from exc

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    validate_required_headers(headers)

    rows: List[Dict[str, Any]] = []
    for idx, row in enumerate(rows_iter, start=2):
        if not any(cell is not None for cell in row):
            continue
        item = {headers[col_idx]: cell for col_idx, cell in enumerate(row) if col_idx < len(headers)}
        item["_xlsx_row_number"] = idx
        rows.append(item)
        if limit > 0 and len(rows) >= limit:
            break

    return headers, rows


def normalize_xlsx_row(row: Dict[str, Any]) -> Dict[str, Any]:
    vid = clean_scalar(row.get("id"))
    project_name = clean_scalar(row.get("Проект"))
    phone = normalize_phone(row.get("Телефон"))
    created_at, created_at_error = parse_created_at(row.get("Создано"))
    subdomain = clean_scalar(row.get("Комментарий"))
    prov_chanel, prov_source = parse_page_parts(project_name)
    phones_raw = [phone] if phone else None

    errors: List[str] = []
    if not vid:
        errors.append("empty_vid")
    if not project_name:
        errors.append("empty_project_name")
    if not phone:
        errors.append("empty_phone")
    if created_at_error:
        errors.append(created_at_error)

    return {
        "xlsx_row_number": row.get("_xlsx_row_number"),
        "raw": row,
        "vid": vid,
        "project_name": project_name,
        "phone": phone,
        "phones_raw": phones_raw,
        "prov_created_at": created_at,
        "subdomain": subdomain,
        "prov_chanel": prov_chanel,
        "prov_source": prov_source,
        "errors": errors,
    }


def chunked(values: List[str], size: int = 500) -> Iterable[List[str]]:
    for start in range(0, len(values), size):
        yield values[start:start + size]


def get_existing_vids(db_sess, vids: List[str]) -> Set[str]:
    if not vids:
        return set()
    existing: Set[str] = set()
    for chunk in chunked(vids, size=500):
        rows = db_sess.execute(
            select(models.ProviderLead.vid).where(models.ProviderLead.vid.in_(chunk))
        ).all()
        for (vid,) in rows:
            if vid is not None:
                existing.add(str(vid))
    return existing


def build_row_preview(normalized_row: Dict[str, Any], note: str) -> Dict[str, Any]:
    return {
        "xlsx_row_number": normalized_row.get("xlsx_row_number"),
        "vid": normalized_row.get("vid"),
        "project_name": normalized_row.get("project_name"),
        "phone": normalized_row.get("phone"),
        "subdomain": normalized_row.get("subdomain"),
        "note": note,
    }


def analyze_rows(db_sess, normalized_rows: List[Dict[str, Any]]) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    report: Dict[str, Any] = {
        "total_rows": len(normalized_rows),
        "valid_rows": 0,
        "duplicates_in_file": 0,
        "duplicates_in_db": 0,
        "new_rows": 0,
        "ready_to_import": 0,
        "matched_projects": 0,
        "not_found_projects": 0,
        "ambiguous_projects": 0,
        "rows_with_errors": 0,
        "errors_breakdown": {},
        "samples": {
            "duplicates_in_file": [],
            "duplicates_in_db": [],
            "not_found_projects": [],
            "ambiguous_projects": [],
            "rows_with_errors": [],
        },
    }

    seen_vids: Set[str] = set()
    candidate_vids = [row["vid"] for row in normalized_rows if row.get("vid")]
    existing_vids = get_existing_vids(db_sess, candidate_vids)
    rows_to_import: List[Dict[str, Any]] = []

    for row in normalized_rows:
        errors = list(row.get("errors") or [])
        vid = row.get("vid")

        if errors:
            report["rows_with_errors"] += 1
            for error_code in errors:
                report["errors_breakdown"][error_code] = report["errors_breakdown"].get(error_code, 0) + 1
            if len(report["samples"]["rows_with_errors"]) < 10:
                report["samples"]["rows_with_errors"].append(build_row_preview(row, ", ".join(errors)))
            continue

        report["valid_rows"] += 1

        if vid in seen_vids:
            report["duplicates_in_file"] += 1
            if len(report["samples"]["duplicates_in_file"]) < 10:
                report["samples"]["duplicates_in_file"].append(build_row_preview(row, "duplicate_vid_in_file"))
            continue
        seen_vids.add(str(vid))

        if str(vid) in existing_vids:
            report["duplicates_in_db"] += 1
            if len(report["samples"]["duplicates_in_db"]) < 10:
                report["samples"]["duplicates_in_db"].append(build_row_preview(row, "duplicate_vid_in_db"))
            continue

        project_id, match_status, matched_projects = crud.resolve_project_by_name_for_provider_lead(
            db_sess,
            str(row["project_name"]),
        )
        row["project_id"] = project_id
        row["project_match_status"] = match_status
        row["matched_project_ids"] = [int(project.id) for project in matched_projects]

        if match_status == "matched":
            report["matched_projects"] += 1
        elif match_status == "not_found":
            report["not_found_projects"] += 1
            if len(report["samples"]["not_found_projects"]) < 10:
                report["samples"]["not_found_projects"].append(build_row_preview(row, "project_not_found"))
        else:
            report["ambiguous_projects"] += 1
            if len(report["samples"]["ambiguous_projects"]) < 10:
                note = f"ambiguous_project_ids={row['matched_project_ids']}"
                report["samples"]["ambiguous_projects"].append(build_row_preview(row, note))

        report["new_rows"] += 1
        rows_to_import.append(row)

    report["ready_to_import"] = len(rows_to_import)
    return report, rows_to_import


def print_report(report: Dict[str, Any], *, verbose: bool = False) -> None:
    print("")
    print("[SUMMARY]")
    print(f"total_rows={report['total_rows']}")
    print(f"valid_rows={report['valid_rows']}")
    print(f"rows_with_errors={report['rows_with_errors']}")
    print(f"duplicates_in_file={report['duplicates_in_file']}")
    print(f"duplicates_in_db={report['duplicates_in_db']}")
    print(f"new_rows={report['new_rows']}")
    print(f"ready_to_import={report['ready_to_import']}")
    print(f"matched_projects={report['matched_projects']}")
    print(f"not_found_projects={report['not_found_projects']}")
    print(f"ambiguous_projects={report['ambiguous_projects']}")

    if report["errors_breakdown"]:
        print("")
        print("[ERRORS_BREAKDOWN]")
        for key in sorted(report["errors_breakdown"].keys()):
            print(f"{key}={report['errors_breakdown'][key]}")

    if not verbose:
        return

    for sample_group in (
        "rows_with_errors",
        "duplicates_in_file",
        "duplicates_in_db",
        "not_found_projects",
        "ambiguous_projects",
    ):
        sample_rows = report["samples"].get(sample_group) or []
        if not sample_rows:
            continue
        print("")
        print(f"[SAMPLE:{sample_group}]")
        for item in sample_rows:
            print(json.dumps(item, ensure_ascii=False, default=str))


def build_provider_lead_model(row: Dict[str, Any]) -> models.ProviderLead:
    return models.ProviderLead(
        vid=str(row["vid"]),
        phone=row.get("phone"),
        phones_raw=row.get("phones_raw"),
        project_name=row.get("project_name"),
        prov_created_at=row.get("prov_created_at"),
        prov_chanel=row.get("prov_chanel"),
        prov_source=row.get("prov_source"),
        subdomain=row.get("subdomain"),
        project_id=row.get("project_id"),
        imported_at=now_msk(),
    )


def insert_provider_leads(db_sess, rows_to_import: List[Dict[str, Any]]) -> int:
    if not rows_to_import:
        return 0
    models_to_insert = [build_provider_lead_model(row) for row in rows_to_import]
    db_sess.add_all(models_to_insert)
    db_sess.commit()
    return len(models_to_insert)


def main() -> None:
    load_dotenv()
    args = parse_args()

    file_path = args.file
    if not os.path.exists(file_path):
        raise SystemExit(f"File not found: {file_path}")

    preview_mode = not args.do_import or args.preview
    _, raw_rows = load_xlsx_rows(file_path, limit=max(0, int(args.limit or 0)))
    normalized_rows = [normalize_xlsx_row(row) for row in raw_rows]

    _, session_local = db_mod.init_engine_and_session(args.db_url)
    db_sess = session_local()

    try:
        report, rows_to_import = analyze_rows(db_sess, normalized_rows)

        if preview_mode and not args.do_import:
            print_report(report, verbose=args.verbose)
            return

        print("[INFO] Preview before import:")
        print_report(report, verbose=args.verbose)
        print("")
        print("[INFO] Starting import...")
        inserted = insert_provider_leads(db_sess, rows_to_import)
        print(f"[OK] Imported rows: {inserted}")

    except SQLAlchemyError as exc:
        db_sess.rollback()
        raise SystemExit(f"Database error: {exc}") from exc
    finally:
        db_sess.close()


if __name__ == "__main__":
    main()
