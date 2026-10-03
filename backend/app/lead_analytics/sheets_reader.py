from __future__ import annotations

import os
import re
import time
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

from dotenv import load_dotenv
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[3]
SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets.readonly"
INVALID_TITLE_CHARS = re.compile(r"[\\/?*\[\]:]")
RETRYABLE_STATUSES = {429, 500, 502, 503, 504}


def parse_spreadsheet_url(url: str) -> tuple[str, int | None]:
    parsed = urlsplit(url.strip())
    match = re.fullmatch(r"/spreadsheets/d/([A-Za-z0-9_-]+)(?:/.*)?", parsed.path)
    if parsed.scheme != "https" or parsed.netloc != "docs.google.com" or not match:
        raise ValueError("Вставьте ссылку на Google Таблицу из docs.google.com")
    gid_value = (parse_qs(parsed.fragment).get("gid") or parse_qs(parsed.query).get("gid") or [None])[0]
    if gid_value is not None and not gid_value.isdigit():
        raise ValueError("В ссылке указан некорректный номер вкладки Google Таблицы")
    return match.group(1), int(gid_value) if gid_value is not None else None


def read_spreadsheet_url(url: str) -> tuple[bytes, str, str | None]:
    spreadsheet_id, gid = parse_spreadsheet_url(url)
    credentials = Credentials.from_service_account_file(_credentials_path(), scopes=[SHEETS_SCOPE])
    service = build("sheets", "v4", credentials=credentials, cache_discovery=False)
    metadata = _execute(service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="properties(title),sheets(properties(sheetId,title))",
    ))
    properties = [sheet["properties"] for sheet in metadata.get("sheets", [])]
    if not properties:
        raise ValueError("В Google Таблице нет вкладок")
    selected_title = next((item["title"] for item in properties if gid is not None and item["sheetId"] == gid), None)
    if gid is not None and selected_title is None:
        raise ValueError(f"Вкладка gid={gid} не найдена в Google Таблице")

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    selected_sheet = None
    active_index = 0
    for item in properties:
        title = str(item["title"])
        sheet_title = _xlsx_sheet_title(title, used_titles)
        worksheet = workbook.create_sheet(sheet_title)
        if title == selected_title:
            selected_sheet = sheet_title
            active_index = workbook.index(worksheet)
        quoted_title = "'" + title.replace("'", "''") + "'"
        values = _execute(service.spreadsheets().values().get(
            spreadsheetId=spreadsheet_id, range=f"{quoted_title}!A:ZZZ",
            valueRenderOption="UNFORMATTED_VALUE", dateTimeRenderOption="SERIAL_NUMBER",
        )).get("values", [])
        column_count = max((len(row) for row in values), default=0)
        if not values or not column_count:
            continue
        end_column = get_column_letter(column_count)
        grid = _execute(service.spreadsheets().get(
            spreadsheetId=spreadsheet_id,
            ranges=[f"{quoted_title}!A1:{end_column}{len(values)}"],
            includeGridData=True,
            fields="sheets(data(rowData(values(effectiveValue,effectiveFormat(numberFormat(type))))))",
        )).get("sheets", [{}])[0].get("data", [{}])[0].get("rowData", [])
        for row_index, row in enumerate(values):
            cells = grid[row_index].get("values", []) if row_index < len(grid) else []
            for column_index in range(column_count):
                cell = cells[column_index] if column_index < len(cells) else {}
                fallback = row[column_index] if column_index < len(row) else None
                _set_cell(worksheet.cell(row_index + 1, column_index + 1), cell, fallback)
    workbook.active = active_index
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), str(metadata.get("properties", {}).get("title") or "Google Таблица"), selected_sheet


def _credentials_path() -> Path:
    load_dotenv(PROJECT_ROOT / ".env")
    value = os.getenv("GOOGLE_CREDENTIALS_FILE", "").strip()
    if not value:
        raise RuntimeError("Не задана переменная GOOGLE_CREDENTIALS_FILE")
    path = Path(value)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    if not path.is_file():
        raise RuntimeError("Не найден файл доступа к Google Таблицам")
    return path


def _xlsx_sheet_title(title: str, used: set[str]) -> str:
    base = INVALID_TITLE_CHARS.sub("-", title).strip(" '")[:31].rstrip() or "Лист"
    candidate, counter = base, 2
    while candidate.casefold() in {item.casefold() for item in used}:
        suffix = f" ({counter})"
        candidate = f"{base[:31 - len(suffix)].rstrip()}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _set_cell(cell: Any, cell_data: dict[str, Any], fallback: Any) -> None:
    value = cell_data.get("effectiveValue", {})
    if "numberValue" in value:
        cell.value = value["numberValue"]
        kind = cell_data.get("effectiveFormat", {}).get("numberFormat", {}).get("type")
        cell.number_format = {"DATE": "yyyy-mm-dd", "DATE_TIME": "yyyy-mm-dd hh:mm:ss", "TIME": "hh:mm:ss", "DURATION": "[h]:mm:ss"}.get(kind, "General")
    elif "stringValue" in value:
        cell.value = value["stringValue"]
    elif "boolValue" in value:
        cell.value = value["boolValue"]
    elif "errorValue" in value:
        cell.value = value["errorValue"].get("type", "#ERROR!")
    else:
        cell.value = fallback if fallback != "" else None


def _execute(request: Any) -> Any:
    for attempt in range(5):
        try:
            return request.execute()
        except HttpError as exc:
            if exc.resp.status not in RETRYABLE_STATUSES or attempt == 4:
                raise
            time.sleep(2**attempt)
    raise RuntimeError("Google API request retry loop ended unexpectedly")
