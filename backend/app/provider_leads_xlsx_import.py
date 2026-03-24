"""
Переиспользуемая логика preview/import лидов из XLSX-выгрузки провайдера.

Модуль используется:
- CLI-утилитой tool_import_provider_leads_from_xlsx.py
- admin endpoint'ами preview/import в FastAPI
"""

from __future__ import annotations

import json
import os
import shutil
import tempfile
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook
from sqlalchemy import select

from . import crud, models
from .time_utils import now_msk


REQUIRED_HEADERS = ["id", "Проект", "Телефон", "Создано", "Комментарий"]
PREVIEW_TTL_SECONDS = int(os.getenv("PROVIDER_LEADS_IMPORT_PREVIEW_TTL_SECONDS", "3600"))
MAX_UPLOAD_SIZE_BYTES = int(os.getenv("PROVIDER_LEADS_IMPORT_MAX_FILE_BYTES", str(10 * 1024 * 1024)))
PREVIEW_STORE_DIR = Path(tempfile.gettempdir()) / "mb_lk_provider_leads_import_previews"
ALLOWED_SUFFIXES = {".xlsx"}


class ProviderLeadsImportError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _ensure_preview_store_dir() -> Path:
    PREVIEW_STORE_DIR.mkdir(parents=True, exist_ok=True)
    return PREVIEW_STORE_DIR


def _preview_dir(preview_id: str) -> Path:
    return _ensure_preview_store_dir() / preview_id


def _preview_state_path(preview_id: str) -> Path:
    return _preview_dir(preview_id) / "preview_state.json"


def _preview_upload_path(preview_id: str, file_name: str) -> Path:
    suffix = Path(file_name).suffix or ".xlsx"
    return _preview_dir(preview_id) / f"source{suffix}"


def _now_utc() -> datetime:
    return datetime.now(timezone.utc)


def _parse_created_at_utc(value: str) -> datetime:
    return datetime.fromisoformat(value)


def cleanup_expired_previews() -> None:
    root = _ensure_preview_store_dir()
    expire_before = _now_utc() - timedelta(seconds=max(60, PREVIEW_TTL_SECONDS))

    for child in root.iterdir():
        if not child.is_dir():
            continue
        state_path = child / "preview_state.json"
        if not state_path.exists():
            shutil.rmtree(child, ignore_errors=True)
            continue
        try:
            payload = json.loads(state_path.read_text(encoding="utf-8"))
            created_at_raw = str(payload.get("created_at") or "").strip()
            created_at = _parse_created_at_utc(created_at_raw)
        except Exception:
            shutil.rmtree(child, ignore_errors=True)
            continue
        if created_at < expire_before:
            shutil.rmtree(child, ignore_errors=True)


def _validate_file_name(file_name: str) -> None:
    suffix = Path(file_name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ProviderLeadsImportError("Поддерживаются только файлы .xlsx", status_code=400)


def _validate_file_size(file_bytes: bytes) -> None:
    if len(file_bytes) > MAX_UPLOAD_SIZE_BYTES:
        raise ProviderLeadsImportError(
            f"Файл слишком большой. Максимум {MAX_UPLOAD_SIZE_BYTES} байт.",
            status_code=400,
        )


def clean_scalar(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.lower() == "none":
        return None
    return text


def normalize_phone(value: Any) -> Optional[str]:
    text = clean_scalar(value)
    if not text:
        return None
    return "".join(text.split())


def parse_created_at(value: Any) -> Tuple[Optional[datetime], Optional[str]]:
    text = clean_scalar(value)
    if not text:
        return None, "empty_created_at"

    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt), None
        except ValueError:
            continue
    return None, f"invalid_created_at:{text}"


def parse_page_parts(page: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    if not page:
        return None, None
    parts = page.split("_")
    prov_chanel = parts[0].strip() if parts else None
    prov_source = None
    if len(parts) >= 3:
        prov_source = "_".join(parts[2:]).strip() or None
    return prov_chanel, prov_source


def validate_required_headers(headers: List[str]) -> None:
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        missing_str = ", ".join(missing)
        raise ProviderLeadsImportError(f"В XLSX нет обязательных колонок: {missing_str}", status_code=400)


def load_xlsx_rows(file_path: str, limit: int = 0) -> Tuple[List[str], List[Dict[str, Any]]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows_iter = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ProviderLeadsImportError("XLSX файл пуст.", status_code=400) from exc

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    validate_required_headers(headers)

    rows: List[Dict[str, Any]] = []
    for idx, row in enumerate(rows_iter, start=2):
        if not any(cell is not None for cell in row):
            continue
        item = {headers[col_idx]: cell for col_idx, cell in enumerate(row) if col_idx < len(headers)}
        item["_xlsx_row_number"] = idx
        rows.append(item)
        if limit > 0 and len(rows) >= limit:
            break

    return headers, rows


def normalize_xlsx_row(row: Dict[str, Any]) -> Dict[str, Any]:
    vid = clean_scalar(row.get("id"))
    project_name = clean_scalar(row.get("Проект"))
    phone = normalize_phone(row.get("Телефон"))
    created_at, created_at_error = parse_created_at(row.get("Создано"))
    subdomain = clean_scalar(row.get("Комментарий"))
    prov_chanel, prov_source = parse_page_parts(project_name)
    phones_raw = [phone] if phone else None

    errors: List[str] = []
    if not vid:
        errors.append("empty_vid")
    if not project_name:
        errors.append("empty_project_name")
    if not phone:
        errors.append("empty_phone")
    if created_at_error:
        errors.append(created_at_error)

    return {
        "xlsx_row_number": row.get("_xlsx_row_number"),
        "raw": row,
        "vid": vid,
        "project_name": project_name,
        "phone": phone,
        "phones_raw": phones_raw,
        "prov_created_at": created_at,
        "subdomain": subdomain,
        "prov_chanel": prov_chanel,
        "prov_source": prov_source,
        "errors": errors,
    }


def chunked(values: List[str], size: int = 500) -> Iterable[List[str]]:
    for start in range(0, len(values), size):
        yield values[start:start + size]


def get_existing_vids(db_sess, vids: List[str]) -> Set[str]:
    if not vids:
        return set()
    existing: Set[str] = set()
    for chunk in chunked(vids, size=500):
        rows = db_sess.execute(
            select(models.ProviderLead.vid).where(models.ProviderLead.vid.in_(chunk))
        ).all()
        for (vid,) in rows:
            if vid is not None:
                existing.add(str(vid))
    return existing


def build_row_preview(normalized_row: Dict[str, Any], note: str) -> Dict[str, Any]:
    return {
        "xlsxRowNumber": normalized_row.get("xlsx_row_number"),
        "vid": normalized_row.get("vid"),
        "projectName": normalized_row.get("project_name"),
        "phone": normalized_row.get("phone"),
        "subdomain": normalized_row.get("subdomain"),
        "note": note,
    }


def analyze_rows(db_sess, normalized_rows: List[Dict[str, Any]]) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    report: Dict[str, Any] = {
        "totalRows": len(normalized_rows),
        "validRows": 0,
        "duplicatesInFile": 0,
        "duplicatesInDb": 0,
        "newRows": 0,
        "readyToImport": 0,
        "matchedProjects": 0,
        "notFoundProjects": 0,
        "ambiguousProjects": 0,
        "rowsWithErrors": 0,
        "errorsBreakdown": {},
        "samples": {
            "duplicatesInFile": [],
            "duplicatesInDb": [],
            "notFoundProjects": [],
            "ambiguousProjects": [],
            "rowsWithErrors": [],
        },
    }

    seen_vids: Set[str] = set()
    candidate_vids = [row["vid"] for row in normalized_rows if row.get("vid")]
    existing_vids = get_existing_vids(db_sess, candidate_vids)
    rows_to_import: List[Dict[str, Any]] = []

    for row in normalized_rows:
        errors = list(row.get("errors") or [])
        vid = row.get("vid")

        if errors:
            report["rowsWithErrors"] += 1
            for error_code in errors:
                report["errorsBreakdown"][error_code] = report["errorsBreakdown"].get(error_code, 0) + 1
            if len(report["samples"]["rowsWithErrors"]) < 10:
                report["samples"]["rowsWithErrors"].append(build_row_preview(row, ", ".join(errors)))
            continue

        report["validRows"] += 1

        if vid in seen_vids:
            report["duplicatesInFile"] += 1
            if len(report["samples"]["duplicatesInFile"]) < 10:
                report["samples"]["duplicatesInFile"].append(build_row_preview(row, "duplicate_vid_in_file"))
            continue
        seen_vids.add(str(vid))

        if str(vid) in existing_vids:
            report["duplicatesInDb"] += 1
            if len(report["samples"]["duplicatesInDb"]) < 10:
                report["samples"]["duplicatesInDb"].append(build_row_preview(row, "duplicate_vid_in_db"))
            continue

        project_id, match_status, matched_projects = crud.resolve_project_by_name_for_provider_lead(
            db_sess,
            str(row["project_name"]),
        )
        row["project_id"] = project_id
        row["project_match_status"] = match_status
        row["matched_project_ids"] = [int(project.id) for project in matched_projects]

        if match_status == "matched":
            report["matchedProjects"] += 1
        elif match_status == "not_found":
            report["notFoundProjects"] += 1
            if len(report["samples"]["notFoundProjects"]) < 10:
                report["samples"]["notFoundProjects"].append(build_row_preview(row, "project_not_found"))
        else:
            report["ambiguousProjects"] += 1
            if len(report["samples"]["ambiguousProjects"]) < 10:
                note = f"ambiguous_project_ids={row['matched_project_ids']}"
                report["samples"]["ambiguousProjects"].append(build_row_preview(row, note))

        report["newRows"] += 1
        rows_to_import.append(row)

    report["readyToImport"] = len(rows_to_import)
    return report, rows_to_import


def build_provider_lead_model(row: Dict[str, Any]) -> models.ProviderLead:
    return models.ProviderLead(
        vid=str(row["vid"]),
        phone=row.get("phone"),
        phones_raw=row.get("phones_raw"),
        project_name=row.get("project_name"),
        prov_created_at=row.get("prov_created_at"),
        prov_chanel=row.get("prov_chanel"),
        prov_source=row.get("prov_source"),
        subdomain=row.get("subdomain"),
        project_id=row.get("project_id"),
        imported_at=now_msk(),
    )


def insert_provider_leads(db_sess, rows_to_import: List[Dict[str, Any]]) -> int:
    if not rows_to_import:
        return 0
    models_to_insert = [build_provider_lead_model(row) for row in rows_to_import]
    db_sess.add_all(models_to_insert)
    db_sess.commit()
    return len(models_to_insert)


def _serialize_import_row(row: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "xlsx_row_number": row.get("xlsx_row_number"),
        "vid": row.get("vid"),
        "project_name": row.get("project_name"),
        "phone": row.get("phone"),
        "phones_raw": row.get("phones_raw"),
        "prov_created_at": row.get("prov_created_at").isoformat() if row.get("prov_created_at") else None,
        "subdomain": row.get("subdomain"),
        "prov_chanel": row.get("prov_chanel"),
        "prov_source": row.get("prov_source"),
        "project_id": row.get("project_id"),
        "project_match_status": row.get("project_match_status"),
        "matched_project_ids": row.get("matched_project_ids") or [],
    }


def _deserialize_import_row(row: Dict[str, Any]) -> Dict[str, Any]:
    prov_created_at = row.get("prov_created_at")
    return {
        "xlsx_row_number": row.get("xlsx_row_number"),
        "vid": row.get("vid"),
        "project_name": row.get("project_name"),
        "phone": row.get("phone"),
        "phones_raw": row.get("phones_raw"),
        "prov_created_at": datetime.fromisoformat(prov_created_at) if prov_created_at else None,
        "subdomain": row.get("subdomain"),
        "prov_chanel": row.get("prov_chanel"),
        "prov_source": row.get("prov_source"),
        "project_id": row.get("project_id"),
        "project_match_status": row.get("project_match_status"),
        "matched_project_ids": row.get("matched_project_ids") or [],
    }


def _save_preview_state(preview_id: str, payload: Dict[str, Any]) -> None:
    preview_dir = _preview_dir(preview_id)
    preview_dir.mkdir(parents=True, exist_ok=True)
    _preview_state_path(preview_id).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _load_preview_state(preview_id: str) -> Dict[str, Any]:
    state_path = _preview_state_path(preview_id)
    if not state_path.exists():
        raise ProviderLeadsImportError("Preview не найден или уже истёк.", status_code=404)
    try:
        return json.loads(state_path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ProviderLeadsImportError("Повреждён preview import-сессии.", status_code=500) from exc


def _delete_preview_state(preview_id: str) -> None:
    shutil.rmtree(_preview_dir(preview_id), ignore_errors=True)


def create_preview(
    db_sess,
    *,
    admin_user_id: int,
    file_name: str,
    file_bytes: bytes,
    limit: int = 0,
) -> Dict[str, Any]:
    cleanup_expired_previews()
    _validate_file_name(file_name)
    _validate_file_size(file_bytes)

    preview_id = uuid.uuid4().hex
    preview_dir = _preview_dir(preview_id)
    preview_dir.mkdir(parents=True, exist_ok=True)
    upload_path = _preview_upload_path(preview_id, file_name)
    upload_path.write_bytes(file_bytes)

    headers, raw_rows = load_xlsx_rows(str(upload_path), limit=max(0, int(limit or 0)))
    normalized_rows = [normalize_xlsx_row(row) for row in raw_rows]
    report, rows_to_import = analyze_rows(db_sess, normalized_rows)

    state_payload = {
        "preview_id": preview_id,
        "admin_user_id": int(admin_user_id),
        "file_name": file_name,
        "created_at": _now_utc().isoformat(),
        "headers": headers,
        "summary": report,
        "rows_to_import": [_serialize_import_row(row) for row in rows_to_import],
    }
    _save_preview_state(preview_id, state_payload)

    return {
        "previewId": preview_id,
        "fileName": file_name,
        **report,
    }


def commit_preview(
    db_sess,
    *,
    preview_id: str,
    admin_user_id: int,
) -> Dict[str, Any]:
    cleanup_expired_previews()
    state = _load_preview_state(preview_id)

    if int(state.get("admin_user_id") or 0) != int(admin_user_id):
        raise ProviderLeadsImportError("Эта preview-сессия принадлежит другому админу.", status_code=403)

    created_at_raw = str(state.get("created_at") or "").strip()
    if not created_at_raw:
        raise ProviderLeadsImportError("Preview невалиден: нет created_at.", status_code=500)
    created_at = _parse_created_at_utc(created_at_raw)
    if created_at < (_now_utc() - timedelta(seconds=max(60, PREVIEW_TTL_SECONDS))):
        _delete_preview_state(preview_id)
        raise ProviderLeadsImportError("Preview истёк. Загрузите файл заново.", status_code=404)

    summary = state.get("summary") or {}
    if int(summary.get("notFoundProjects") or 0) > 0 or int(summary.get("ambiguousProjects") or 0) > 0:
        raise ProviderLeadsImportError(
            "Импорт заблокирован: есть строки без однозначной привязки к проекту.",
            status_code=400,
        )

    stored_rows = state.get("rows_to_import") or []
    rows_to_import = [_deserialize_import_row(row) for row in stored_rows]
    existing_vids = get_existing_vids(db_sess, [str(row.get("vid")) for row in rows_to_import if row.get("vid")])
    fresh_rows = [row for row in rows_to_import if str(row.get("vid")) not in existing_vids]

    inserted_rows = insert_provider_leads(db_sess, fresh_rows)
    skipped_duplicates_in_db = len(rows_to_import) - len(fresh_rows)

    _delete_preview_state(preview_id)
    return {
        "previewId": preview_id,
        "fileName": str(state.get("file_name") or ""),
        "insertedRows": inserted_rows,
        "skippedDuplicatesInDb": skipped_duplicates_in_db,
    }


def build_preview_report_payload(file_path: str, db_url: str, headers: List[str], limit: int, summary: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "file": os.path.abspath(file_path),
        "db_url": db_url,
        "headers": headers,
        "mode": "preview",
        "limit": max(0, int(limit or 0)),
        "summary": summary,
        "generated_at": datetime.now().isoformat(),
    }


def build_import_report_payload(file_path: str, db_url: str, headers: List[str], limit: int, summary: Dict[str, Any], inserted_rows: int) -> Dict[str, Any]:
    payload = build_preview_report_payload(file_path, db_url, headers, limit, summary)
    payload["mode"] = "import"
    payload["inserted_rows"] = inserted_rows
    return payload
